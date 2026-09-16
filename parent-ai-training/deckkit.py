# -*- coding: utf-8 -*-
"""학부모 연수용 슬라이드 제작 도구 (공통 레이아웃/색상/도형 헬퍼)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

FONT = "맑은 고딕"

# 흰 배경 / 진한 회색 제목 / 파랑·청록 강조. 신호등 활동에만 별도 색.
C = {
    "bg":       RGBColor(0xFF, 0xFF, 0xFF),
    "ink":      RGBColor(0x1F, 0x29, 0x33),  # 제목
    "body":     RGBColor(0x3E, 0x4C, 0x59),  # 본문
    "mute":     RGBColor(0x7B, 0x88, 0x94),  # 보조/출처
    "blue":     RGBColor(0x1D, 0x4E, 0xD8),  # 강조 1
    "teal":     RGBColor(0x0E, 0x7C, 0x86),  # 강조 2
    "tint":     RGBColor(0xF3, 0xF6, 0xFA),  # 카드 배경
    "tintblue": RGBColor(0xE8, 0xEF, 0xFC),
    "tintteal": RGBColor(0xE4, 0xF2, 0xF3),
    "line":     RGBColor(0xD9, 0xE0, 0xE8),
    "white":    RGBColor(0xFF, 0xFF, 0xFF),
    # 신호등 전용
    "green":    RGBColor(0x15, 0x80, 0x3D),
    "amber":    RGBColor(0xB4, 0x53, 0x09),
    "red":      RGBColor(0xB9, 0x1C, 0x1C),
    "greenbg":  RGBColor(0xE7, 0xF5, 0xEC),
    "amberbg":  RGBColor(0xFD, 0xF1, 0xE3),
    "redbg":    RGBColor(0xFB, 0xEA, 0xEA),
}

W = 13.333
H = 7.5
ML = 0.85          # 좌우 여백
CW = W - 2 * ML    # 콘텐츠 폭


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _ea(run, font=FONT):
    """한글(동아시아) 글꼴을 별도로 지정한다."""
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def txt(slide, l, t, w, h, text, size=18, bold=False, color="body",
        align=PP_ALIGN.LEFT, spacing=1.15, anchor=MSO_ANCHOR.TOP,
        space_after=0, font=FONT, wrap=True):
    """여러 줄 텍스트 상자. text는 str 또는 list[str]."""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, (list, tuple)) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = C[color] if isinstance(color, str) else color
        _ea(r, font)
    return box


def rich(slide, l, t, w, h, parts, size=18, align=PP_ALIGN.LEFT,
         spacing=1.15, anchor=MSO_ANCHOR.TOP, space_after=0):
    """한 문단 안에서 색·굵기를 섞는다. parts = [(text, color, bold), ...] 리스트의 리스트."""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = parts if isinstance(parts[0], list) else [parts]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        for seg in para:
            s_text, s_color, s_bold = (list(seg) + [None, None])[:3]
            r = p.add_run()
            r.text = s_text
            r.font.size = Pt(size)
            r.font.bold = bool(s_bold)
            r.font.name = FONT
            r.font.color.rgb = C[s_color or "body"] if isinstance(s_color or "body", str) else s_color
            _ea(r)
    return box


def rrect(slide, l, t, w, h, fill="tint", line=None, radius=0.06, shadow=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(l), Inches(t), Inches(w), Inches(h))
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = C[fill] if isinstance(fill, str) else fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = C[line] if isinstance(line, str) else line
        sh.line.width = Pt(1.25)
    if not shadow:
        sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def rect(slide, l, t, w, h, fill="tint", line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = C[fill] if isinstance(fill, str) else fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = C[line] if isinstance(line, str) else line
        sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh


def oval(slide, l, t, w, h, fill="tintblue", line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = C[fill] if isinstance(fill, str) else fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = C[line] if isinstance(line, str) else line
        sh.line.width = Pt(1.5)
    sh.shadow.inherit = False
    return sh


def chevron(slide, l, t, w, h, fill="tintblue"):
    sh = slide.shapes.add_shape(MSO_SHAPE.CHEVRON,
                                Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = C[fill] if isinstance(fill, str) else fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def arrow(slide, l, t, w, h=0.34, fill="blue"):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = C[fill] if isinstance(fill, str) else fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def title(slide, text, kicker=None, sub=None, top=0.62):
    """상단 제목 블록. 얇은 강조선 + 제목(+ 보조 문장)."""
    y = top
    if kicker:
        txt(slide, ML, y - 0.30, CW, 0.28, kicker, size=13.5, bold=True, color="teal")
    rect(slide, ML, y, 0.62, 0.055, fill="blue")
    y += 0.20
    tb = txt(slide, ML, y, CW, 0.75, text, size=32, bold=True, color="ink", spacing=1.05)
    y += 0.78
    if sub:
        txt(slide, ML, y, CW, 0.42, sub, size=16, color="mute", spacing=1.2)
        y += 0.50
    return y


def footer(slide, page, source=None, total=30):
    rect(slide, 0, H - 0.52, W, 0.012, fill="line")
    if source:
        txt(slide, ML, H - 0.40, CW - 1.2, 0.28, source, size=9.5, color="mute")
    txt(slide, W - ML - 1.0, H - 0.40, 1.0, 0.28, f"{page} / {total}",
        size=9.5, color="mute", align=PP_ALIGN.RIGHT)


def notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(11)
            r.font.name = FONT
            _ea(r)


def card(slide, l, t, w, h, head, body=None, fill="tint", line="line",
         accent=None, head_size=17, body_size=13.5, align=PP_ALIGN.LEFT, pad=0.28):
    """머리글 + 본문(여러 줄)을 담은 카드."""
    rrect(slide, l, t, w, h, fill=fill, line=line)
    y = t + pad
    if accent:
        rect(slide, l + pad, y, 0.42, 0.05, fill=accent)
        y += 0.18
    txt(slide, l + pad, y, w - 2 * pad, 0.42, head, size=head_size, bold=True,
        color="ink", align=align, spacing=1.12)
    y += 0.30 + (head_size - 17) * 0.012
    if body:
        lines = body if isinstance(body, (list, tuple)) else [body]
        txt(slide, l + pad, y + 0.16, w - 2 * pad, h - (y - t) - pad,
            lines, size=body_size, color="body", align=align, spacing=1.32, space_after=4)
    return y


def numbadge(slide, l, t, d, n, fill="blue", color="white", size=15):
    oval(slide, l, t, d, d, fill=fill)
    txt(slide, l, t + d / 2 - 0.14, d, 0.3, str(n), size=size, bold=True,
        color=color, align=PP_ALIGN.CENTER)
